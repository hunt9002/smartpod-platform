from pptx import Presentation

def create_ppt(data):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])

    slide.shapes.title.text = "Smart POD 설계 결과"
    slide.placeholders[1].text = f"""
    Total Power: {data['Total_kW']} kW
    UPS: {data['UPS_kVA']} kVA
    Cooling: {data['Cooling_RT']} RT
    Flow: {data['Flow_LPM']} LPM
    """

    prs.save("SmartPOD_Proposal.pptx")
